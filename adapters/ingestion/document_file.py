#!/usr/bin/env python3
"""
Conector de ingesta "document_file" -- primero de la familia de documentos
(docs/adr/0015, docs/design/plan-ingesta-documentos.md, Fase A). Lee un documento de
texto ya volcado a disco -- `.docx`, `.txt`, `.md` -- y lo envuelve como RawCapture
(adapters/ingestion/CONTRACT.md), igual que meeting_file.py pero para esta fuente.

A diferencia de meeting_file.py, esta familia sigue el contrato con dos diferencias
deliberadas (docs/adr/0015):

- Nunca persiste el crudo original -- ni este modulo ni el codigo que lo orquesta
  llaman a lib/ingestion.py::save_capture. El archivo se lee, se extrae su texto, y
  se descarta.
- Puede levantar `IngestionExtractionError` (regla 6 del contrato): la fuente esta
  disponible pero su contenido no se puede extraer de forma legible (un `.docx`
  corrupto, un formato que la libreria no reconoce). Distinta de
  `IngestionProviderError` (fuente inalcanzable -- el archivo ni se puede abrir).

Este modulo NO destila nada -- solo trae el contenido crudo con un locator estable.
"""
from __future__ import annotations

import hashlib
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

_TEXT_EXTENSIONS = {".txt", ".md"}
_DOCX_EXTENSIONS = {".docx"}


class IngestionProviderError(RuntimeError):
    """La fuente no se pudo leer -- ausencia explicita (principio 5, regla 3 del
    contrato), nunca un RawCapture vacio o silenciosamente salteado."""


class IngestionExtractionError(RuntimeError):
    """La fuente esta disponible pero su contenido no se puede extraer de forma
    legible (regla 6 del contrato, docs/adr/0015) -- nunca un RawCapture vacio,
    parcial o degradado en silencio."""


def _extract_docx_text(path: Path) -> str:
    try:
        import docx  # python-docx -- ver requirements.txt
    except ImportError as exc:  # pragma: no cover -- dependencia declarada en requirements.txt
        raise IngestionExtractionError(
            "falta la dependencia 'python-docx' para leer archivos .docx (ver requirements.txt)"
        ) from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001 -- python-docx no expone una jerarquia de excepciones propia
        raise IngestionExtractionError(f"el archivo .docx esta corrupto o no se pudo abrir: {path} ({exc})") from exc

    paragraphs = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append("\t".join(cell.text for cell in row.cells))

    return "\n".join(paragraphs)


def fetch_raw(path: str | Path, capture_id: str | None = None) -> dict[str, Any]:
    """Lee el documento en 'path' y lo envuelve como RawCapture
    (adapters/ingestion/CONTRACT.md). 'capture_id' por default es el nombre del
    archivo sin extension -- pasarlo explicito si dos documentos distintos podrian
    compartir nombre de archivo."""
    path = Path(path)
    if not path.is_file():
        raise IngestionProviderError(f"no existe el archivo: {path}")

    suffix = path.suffix.lower()
    if suffix in _TEXT_EXTENSIONS:
        try:
            raw_text = path.read_text(encoding="utf-8")
        except UnicodeDecodeError as exc:
            raise IngestionExtractionError(f"el archivo no es texto UTF-8 legible: {path} ({exc})") from exc
    elif suffix in _DOCX_EXTENSIONS:
        raw_text = _extract_docx_text(path)
    else:
        raise IngestionExtractionError(
            f"extension no soportada por document_file.py: {suffix!r} (soportadas: "
            f"{sorted(_TEXT_EXTENSIONS | _DOCX_EXTENSIONS)})"
        )

    if not raw_text.strip():
        raise IngestionExtractionError(f"no se pudo extraer texto del documento (quedo vacio): {path}")

    return {
        "capture_id": capture_id or path.stem,
        "source": "document",
        "locator": str(path),
        "captured_at": datetime.now(timezone.utc).isoformat(),
        "raw_text": raw_text,
        "content_hash": hashlib.sha256(raw_text.encode("utf-8")).hexdigest(),
    }
