#!/usr/bin/env python3
"""
Conector de ingesta "spreadsheet_file" (docs/adr/0015, docs/design/plan-ingesta-documentos.md
seccion 4.2, Fase C). Lee una hoja de calculo ya volcada a disco -- `.xlsx` o `.csv`
-- y la serializa como texto tabular (filas separadas por salto de linea, columnas
por tab), envolviendola como RawCapture (adapters/ingestion/CONTRACT.md). `source`:
`spreadsheet` (enum nuevo, docs/adr/0015 seccion 4.4).

Mismas dos diferencias que el resto de la familia de documentos (docs/adr/0015):

- Nunca persiste el crudo original (nunca llama a lib/ingestion.py::save_capture).
- Puede levantar `IngestionExtractionError` (regla 6 del contrato) si el archivo
  esta disponible pero no se puede leer de forma legible: `.xlsx` corrupto,
  protegido/encriptado, o sin ninguna hoja con contenido.

Convencion de locator para citar una celda o rango especifico como evidencia (no
para el `locator` del RawCapture en si, que sigue siendo el path al archivo
completo): `<archivo>#<hoja>!<rango>` (ej. `plan.xlsx#Presupuesto!A1:C4`), o
`<archivo>#<rango>` para un CSV sin hojas -- ver adapters/ingestion/CONTRACT.md.
"""
from __future__ import annotations

import csv
import hashlib
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


class IngestionProviderError(RuntimeError):
    """La fuente no se pudo leer -- ausencia explicita (principio 5, regla 3 del
    contrato), nunca un RawCapture vacio o silenciosamente salteado."""


class IngestionExtractionError(RuntimeError):
    """La fuente esta disponible pero su contenido no se puede extraer de forma
    legible (regla 6 del contrato, docs/adr/0015) -- nunca un RawCapture vacio,
    parcial o degradado en silencio."""


def _row_to_line(row: tuple) -> str | None:
    if all(cell is None or str(cell).strip() == "" for cell in row):
        return None
    return "\t".join("" if cell is None else str(cell) for cell in row)


def _extract_csv(path: Path) -> str:
    try:
        with path.open(newline="", encoding="utf-8") as f:
            reader = csv.reader(f)
            lines = [line for line in (_row_to_line(tuple(row)) for row in reader) if line is not None]
    except UnicodeDecodeError as exc:
        raise IngestionExtractionError(f"el CSV no es texto UTF-8 legible: {path} ({exc})") from exc
    except csv.Error as exc:
        raise IngestionExtractionError(f"el CSV esta mal formado: {path} ({exc})") from exc

    if not lines:
        raise IngestionExtractionError(f"el CSV no tiene ninguna fila con contenido: {path}")
    return "\n".join(lines)


def _extract_xlsx(path: Path) -> tuple[str, list[str]]:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover -- dependencia declarada en requirements.txt
        raise IngestionExtractionError(
            "falta la dependencia 'openpyxl' para leer archivos .xlsx (ver requirements.txt)"
        ) from exc

    try:
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 -- openpyxl no expone una jerarquia de excepciones cerrada
        raise IngestionExtractionError(
            f"el archivo .xlsx esta corrupto, protegido o no se pudo abrir: {path} ({exc})"
        ) from exc

    sheet_blocks: list[str] = []
    extraction_notes: list[str] = []
    for sheet_name in workbook.sheetnames:
        try:
            sheet = workbook[sheet_name]
            lines = [line for line in (_row_to_line(row) for row in sheet.iter_rows(values_only=True)) if line is not None]
        except Exception as exc:  # noqa: BLE001
            extraction_notes.append(f"hoja {sheet_name!r}: no se pudo leer ({exc})")
            continue
        if not lines:
            extraction_notes.append(f"hoja {sheet_name!r}: sin contenido")
            continue
        sheet_blocks.append(f"=== Hoja: {sheet_name} ===\n" + "\n".join(lines))

    return "\n\n".join(sheet_blocks), extraction_notes


def fetch_raw(path: str | Path, capture_id: str | None = None) -> dict[str, Any]:
    """Lee la hoja de calculo en 'path' y la envuelve como RawCapture
    (adapters/ingestion/CONTRACT.md). 'capture_id' por default es el nombre del
    archivo sin extension -- pasarlo explicito si dos archivos distintos podrian
    compartir nombre."""
    path = Path(path)
    if not path.is_file():
        raise IngestionProviderError(f"no existe el archivo: {path}")

    suffix = path.suffix.lower()
    extraction_notes: list[str] = []
    if suffix == ".csv":
        raw_text = _extract_csv(path)
    elif suffix == ".xlsx":
        raw_text, extraction_notes = _extract_xlsx(path)
        if not raw_text.strip():
            raise IngestionExtractionError(
                f"no se pudo extraer contenido de ninguna hoja del archivo: {path} "
                f"(detalle: {'; '.join(extraction_notes) if extraction_notes else 'sin hojas'})"
            )
    else:
        raise IngestionExtractionError(
            f"extension no soportada por spreadsheet_file.py: {suffix!r} (soportadas: ['.csv', '.xlsx'])"
        )

    capture: dict[str, Any] = {
        "capture_id": capture_id or path.stem,
        "source": "spreadsheet",
        "locator": str(path),
        "captured_at": datetime.now(timezone.utc).isoformat(),
        "raw_text": raw_text,
        "content_hash": hashlib.sha256(raw_text.encode("utf-8")).hexdigest(),
    }
    if extraction_notes:
        capture["extraction_notes"] = extraction_notes
    return capture
