#!/usr/bin/env python3
"""
Prueba hermetica de los cuatro conectores nuevos de la familia de documentos
(docs/adr/0015, docs/design/plan-ingesta-documentos.md, Fases A-D):
document_file.py, pdf_file.py, spreadsheet_file.py, image_file.py.

Nunca toca una fuente externa real -- construye sus propios archivos descartables
(.docx/.pdf/.xlsx/.csv/.png) en un directorio temporal con las mismas librerias que
usan los conectores (python-docx, reportlab, openpyxl, Pillow), corre fetch_raw()
contra ellos, y descarta el directorio al final. Mismo patron que
tests/test-audit.py arma un Context Base y un repo de codigo descartables en vez de
fixtures estaticas -- para estos conectores, generar el binario en el momento es mas
simple y explicito que comitear archivos binarios al repo.

Cubre, para cada conector: extraccion exitosa (source correcto, texto correcto),
fuente inalcanzable (IngestionProviderError), contenido no extraible
(IngestionExtractionError), y las dos excepciones documentadas en
adapters/ingestion/CONTRACT.md: falla parcial declarada via `extraction_notes` (PDF
con una pagina sin texto, xlsx con una hoja vacia) y la imagen sin texto reconocido
(resultado valido, no un error). Tambien confirma la regla central del ADR: ninguno
de los cuatro conectores persiste su captura (no importan ni llaman a
lib/ingestion.py::save_capture).
"""
from __future__ import annotations

import shutil
import sys
import tempfile
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))

from adapters.ingestion import document_file, image_file, pdf_file, spreadsheet_file  # noqa: E402

failures: list[str] = []


def check(description: str, condition: bool) -> None:
    if condition:
        print(f"PASS: {description}")
    else:
        print(f"FAIL: {description}")
        failures.append(description)


def _module_source_has_no_save_capture(module) -> bool:
    """Confirma en el codigo fuente que el conector nunca INVOCA save_capture
    (docs/adr/0015) -- solo se prohibe la llamada real ("save_capture("), no la
    mencion en un comentario/docstring que referencia la regla del ADR."""
    source = Path(module.__file__).read_text(encoding="utf-8")
    return "save_capture(" not in source


def main() -> int:
    tmp_dir = Path(tempfile.mkdtemp(prefix="metis-doc-ingestion-"))
    try:
        # ============================================================
        # Fase A: document_file.py (.docx / .txt / .md)
        # ============================================================
        import docx

        (tmp_dir / "notas.md").write_text("# PRD\n\nContenido de un PRD de prueba.\n", encoding="utf-8")
        cap_md = document_file.fetch_raw(tmp_dir / "notas.md")
        check("document_file: .md exitoso trae source=document", cap_md["source"] == "document")
        check("document_file: .md exitoso trae el texto completo", "Contenido de un PRD de prueba." in cap_md["raw_text"])
        check("document_file: content_hash presente", bool(cap_md.get("content_hash")))

        doc = docx.Document()
        doc.add_paragraph("Primer parrafo del docx de prueba.")
        doc.add_paragraph("Segundo parrafo, distinto.")
        doc.save(str(tmp_dir / "prd.docx"))
        cap_docx = document_file.fetch_raw(tmp_dir / "prd.docx")
        check("document_file: .docx exitoso trae source=document", cap_docx["source"] == "document")
        check("document_file: .docx exitoso extrae ambos parrafos", "Primer parrafo" in cap_docx["raw_text"] and "Segundo parrafo" in cap_docx["raw_text"])

        try:
            document_file.fetch_raw(tmp_dir / "no-existe.docx")
            check("document_file: archivo inexistente levanta IngestionProviderError", False)
        except document_file.IngestionProviderError:
            check("document_file: archivo inexistente levanta IngestionProviderError", True)

        (tmp_dir / "corrupto.docx").write_bytes(b"esto no es un zip valido")
        try:
            document_file.fetch_raw(tmp_dir / "corrupto.docx")
            check("document_file: .docx corrupto levanta IngestionExtractionError", False)
        except document_file.IngestionExtractionError:
            check("document_file: .docx corrupto levanta IngestionExtractionError", True)

        check("document_file: nunca llama a save_capture (docs/adr/0015)", _module_source_has_no_save_capture(document_file))

        # ============================================================
        # Fase B: pdf_file.py
        # ============================================================
        from reportlab.pdfgen import canvas
        from pypdf import PdfReader, PdfWriter

        pdf_path = tmp_dir / "doc.pdf"
        c = canvas.Canvas(str(pdf_path))
        c.drawString(100, 750, "Contenido real de la pagina uno.")
        c.showPage()
        c.drawString(100, 750, "Contenido real de la pagina dos.")
        c.showPage()
        c.save()
        cap_pdf = pdf_file.fetch_raw(pdf_path)
        check("pdf_file: exitoso trae source=document", cap_pdf["source"] == "document")
        check("pdf_file: extrae texto de ambas paginas", "pagina uno" in cap_pdf["raw_text"] and "pagina dos" in cap_pdf["raw_text"])
        check("pdf_file: sin extraction_notes cuando las dos paginas tienen texto", not cap_pdf.get("extraction_notes"))

        # PDF con una pagina en blanco (sin texto) + una con texto -> falla parcial explicita
        mixed_path = tmp_dir / "mixto.pdf"
        c2 = canvas.Canvas(str(mixed_path))
        c2.drawString(100, 750, "Unica pagina con texto real.")
        c2.showPage()
        c2.showPage()  # segunda pagina en blanco, sin drawString
        c2.save()
        cap_mixed = pdf_file.fetch_raw(mixed_path)
        check("pdf_file: falla parcial (1 de 2 paginas sin texto) no lanza excepcion", "Unica pagina con texto real." in cap_mixed["raw_text"])
        check("pdf_file: falla parcial queda declarada en extraction_notes", bool(cap_mixed.get("extraction_notes")) and any("pagina 2" in n for n in cap_mixed["extraction_notes"]))

        blank_path = tmp_dir / "blank.pdf"
        c3 = canvas.Canvas(str(blank_path))
        c3.showPage()
        c3.save()
        try:
            pdf_file.fetch_raw(blank_path)
            check("pdf_file: PDF totalmente sin texto levanta IngestionExtractionError", False)
        except pdf_file.IngestionExtractionError:
            check("pdf_file: PDF totalmente sin texto levanta IngestionExtractionError", True)

        writer = PdfWriter()
        for p in PdfReader(str(pdf_path)).pages:
            writer.add_page(p)
        writer.encrypt(user_password="secreto123")
        enc_path = tmp_dir / "enc.pdf"
        with open(enc_path, "wb") as f:
            writer.write(f)
        try:
            pdf_file.fetch_raw(enc_path)
            check("pdf_file: PDF protegido con contraseña levanta IngestionExtractionError", False)
        except pdf_file.IngestionExtractionError:
            check("pdf_file: PDF protegido con contraseña levanta IngestionExtractionError", True)

        try:
            pdf_file.fetch_raw(tmp_dir / "no-existe.pdf")
            check("pdf_file: archivo inexistente levanta IngestionProviderError", False)
        except pdf_file.IngestionProviderError:
            check("pdf_file: archivo inexistente levanta IngestionProviderError", True)

        check("pdf_file: nunca llama a save_capture (docs/adr/0015)", _module_source_has_no_save_capture(pdf_file))

        # ============================================================
        # Fase C: spreadsheet_file.py (.xlsx / .csv)
        # ============================================================
        import openpyxl

        csv_path = tmp_dir / "datos.csv"
        csv_path.write_text("nombre,edad\nAna,30\nLuis,25\n", encoding="utf-8")
        cap_csv = spreadsheet_file.fetch_raw(csv_path)
        check("spreadsheet_file: .csv exitoso trae source=spreadsheet (enum nuevo)", cap_csv["source"] == "spreadsheet")
        check("spreadsheet_file: .csv serializado como texto tabular (tabs)", "nombre\tedad" in cap_csv["raw_text"] and "Ana\t30" in cap_csv["raw_text"])

        xlsx_path = tmp_dir / "plan.xlsx"
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Presupuesto"
        ws1.append(["Item", "Monto"])
        ws1.append(["Hosting", 100])
        ws2 = wb.create_sheet("Vacia")  # hoja sin contenido -> falla parcial
        wb.save(str(xlsx_path))
        cap_xlsx = spreadsheet_file.fetch_raw(xlsx_path)
        check("spreadsheet_file: .xlsx exitoso trae source=spreadsheet", cap_xlsx["source"] == "spreadsheet")
        check("spreadsheet_file: .xlsx serializa la hoja con contenido", "Presupuesto" in cap_xlsx["raw_text"] and "Hosting\t100" in cap_xlsx["raw_text"])
        check("spreadsheet_file: hoja vacia queda declarada en extraction_notes, no rompe la corrida", bool(cap_xlsx.get("extraction_notes")) and any("Vacia" in n for n in cap_xlsx["extraction_notes"]))

        (tmp_dir / "vacio.csv").write_text("", encoding="utf-8")
        try:
            spreadsheet_file.fetch_raw(tmp_dir / "vacio.csv")
            check("spreadsheet_file: .csv vacio levanta IngestionExtractionError", False)
        except spreadsheet_file.IngestionExtractionError:
            check("spreadsheet_file: .csv vacio levanta IngestionExtractionError", True)

        (tmp_dir / "corrupto.xlsx").write_bytes(b"esto no es un zip valido")
        try:
            spreadsheet_file.fetch_raw(tmp_dir / "corrupto.xlsx")
            check("spreadsheet_file: .xlsx corrupto levanta IngestionExtractionError", False)
        except spreadsheet_file.IngestionExtractionError:
            check("spreadsheet_file: .xlsx corrupto levanta IngestionExtractionError", True)

        try:
            spreadsheet_file.fetch_raw(tmp_dir / "no-existe.csv")
            check("spreadsheet_file: archivo inexistente levanta IngestionProviderError", False)
        except spreadsheet_file.IngestionProviderError:
            check("spreadsheet_file: archivo inexistente levanta IngestionProviderError", True)

        check("spreadsheet_file: nunca llama a save_capture (docs/adr/0015)", _module_source_has_no_save_capture(spreadsheet_file))

        # ============================================================
        # Fase D: image_file.py (OCR)
        # ============================================================
        from PIL import Image, ImageDraw

        img = Image.new("RGB", (500, 120), color="white")
        draw = ImageDraw.Draw(img)
        draw.text((10, 45), "TEXTO RECONOCIBLE DE PRUEBA", fill="black")
        img_path = tmp_dir / "captura.png"
        img.save(img_path)
        cap_img = image_file.fetch_raw(img_path)
        check("image_file: exitoso trae source=image (enum nuevo)", cap_img["source"] == "image")
        check("image_file: OCR reconoce el texto embebido", "RECONOCIBLE" in cap_img["raw_text"].upper())
        check("image_file: sin extraction_notes cuando si hay texto", not cap_img.get("extraction_notes"))

        blank_img = Image.new("RGB", (200, 200), color="white")
        blank_img_path = tmp_dir / "blanco.png"
        blank_img.save(blank_img_path)
        cap_blank = image_file.fetch_raw(blank_img_path)
        check("image_file: imagen sin texto NO levanta excepcion (resultado valido)", cap_blank["raw_text"] == "")
        check("image_file: imagen sin texto lo declara en extraction_notes", bool(cap_blank.get("extraction_notes")))

        try:
            image_file.fetch_raw(tmp_dir / "no-existe.png")
            check("image_file: archivo inexistente levanta IngestionProviderError", False)
        except image_file.IngestionProviderError:
            check("image_file: archivo inexistente levanta IngestionProviderError", True)

        (tmp_dir / "corrupta.png").write_bytes(b"esto no es una imagen valida")
        try:
            image_file.fetch_raw(tmp_dir / "corrupta.png")
            check("image_file: imagen corrupta levanta IngestionExtractionError", False)
        except image_file.IngestionExtractionError:
            check("image_file: imagen corrupta levanta IngestionExtractionError", True)

        check("image_file: nunca llama a save_capture (docs/adr/0015)", _module_source_has_no_save_capture(image_file))

        # ============================================================
        # Vocabulario compartido: los 4 conectores nuevos usan valores de
        # evidence.source que ya existen en el enum de los schemas (docs/adr/0015 4.4)
        # ============================================================
        import json

        schema = json.loads((REPO_ROOT / "schemas" / "requirement.schema.json").read_text(encoding="utf-8"))
        source_enum = schema["properties"]["evidence"]["items"]["properties"]["source"]["enum"]
        check("schema: 'spreadsheet' esta en el enum de evidence.source", "spreadsheet" in source_enum)
        check("schema: 'image' esta en el enum de evidence.source", "image" in source_enum)
        check("schema: 'document' (ya existente) sigue en el enum de evidence.source", "document" in source_enum)

    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)

    print()
    if failures:
        print(f"FALLO: {len(failures)} check(s) no pasaron.")
        return 1
    print("OK: todos los checks pasaron.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
